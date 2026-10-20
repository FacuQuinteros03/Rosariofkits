"""Migración de los libros de Sheets a la base del ERP (diseño 6.1-6.3, plan etapa 3).

Lee los dos libros con permiso de SOLO LECTURA (no puede escribir ni una celda), arma
todo en memoria, lo carga en la base en una sola transacción y compara contra los
libros. Escribe un informe en herramientas/migracion/informes/ (no se commitea: tiene
nombres de clientes y montos, y el repo de la web es público).

    python migrar_erp.py --revisar                  # carga en la base LOCAL y deshace: solo el informe
    python migrar_erp.py --ensayo                   # carga en la base LOCAL y confirma
    python migrar_erp.py --ensayo --destino real    # carga en el proyecto de Supabase (borra lo que haya)
    python migrar_erp.py --final  --destino real    # el corte: lo mismo, con otra confirmación

De dónde sale cada cosa:
  * Libro_RosarioFkits (el nuevo): la base. Productos con modelo y talle separados,
    pedidos con su costo, clientes, ventas hasta V-0026, pagos, caja y ajustes. Ya no
    se carga (Facu, 2026-10-13: "no se entiende al tener tantas cosas").
  * Stock_RosarioFkits (el viejo): el que lee la web y donde se siguen cargando las
    ventas. Manda en nombre, precio, Nuevo y Precio Antes, y aporta las ventas que se
    cargaron después de la unificación (no tiene pagos: se suponen cobradas al vender
    y quedan para que Facu las confirme).
"""

import argparse
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

VIEJO = "1pmjQafZdaLugqD6xbGaqQgZYz6C696QVk5Z4r05w5W4"
NUEVO = "1GuWiUJ5f8jQeFoBlQXSeh_q18_fnshX9VSy-XY5YlJI"
CREDENCIALES = Path(r"C:\Users\facuq\Desktop\RosarioFkits\mcp-sheets\credenciales.json")
AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
LOCAL = "postgresql://postgres:postgres@127.0.0.1:54322/postgres"

TALLES = ["XS", "S", "M", "L", "XL", "XXL", "XXXL", "ÚNICO"]
TALLE_EQUIVALENTE = {"2XL": "XXL", "3XL": "XXXL", "UNICO": "ÚNICO", "U": "ÚNICO"}
METODOS = {"transferencia": "Transferencia", "efectivo": "Efectivo",
           "mercado pago": "Mercado Pago", "mercadopago": "Mercado Pago", "mp": "Mercado Pago"}
ESTADOS = {"reservada": "Reservada", "senada": "Señada", "pagada": "Pagada",
           "entregada": "Entregada", "cancelada": "Cancelada"}
TIPOS_PAGO = {"total": "Total", "sena": "Seña", "saldo": "Saldo", "devolucion": "Devolucion",
              "parcial": "Saldo"}
ESTADOS_PEDIDO = {"encargado": "Encargado", "en transito": "En transito",
                  "recibido": "Recibido", "cancelado": "Cancelado"}
CATEGORIAS_CAJA = ["Venta", "Pago proveedor", "Envio China", "Envio cliente", "Publicidad",
                   "Retiro", "Aporte", "Otro"]

# Correcciones que no están (o no pueden estar) en los libros. Cada una con quién y cuándo.
CORRECCIONES = {
    # Facu, 2026-10-13: pidió 4 Real Madrid Alternativa XL, llegaron 3; el proveedor
    # reemplazó la que faltaba con una Barcelona 25/26 Suplente L (esa ya está en AJUSTES).
    "recibidas": {("P-002", 3020): 3},
    # El libro nuevo dice "En transito", pero ya llegó (Facu, 13/10: "llegó el otro día, no
    # recuerdo la fecha"). Ya estaba en la web el 06/10 y la primera venta es del 08/10.
    "pedido_recibido": {"P-005": date(2026, 10, 6)},
    # Facu, 13/10: el envío de Diez Store fue $17.000 (el libro decía $20.000 estimado).
    "costo_envio": {"P-005": 17000},
    # Facu, 13/10: los pedidos de agosto y el de Diez Store se pagaron con plata propia,
    # antes de que existiera la caja (o por fuera de ella). Cada uno entra como un aporte
    # y sale como pago al proveedor por el mismo monto, el mismo día: la caja no cambia y
    # el pedido deja de figurar como deuda.
    "pedidos_pagados_con_aporte": {
        "P-001": "Facu, 200 USD",
        "P-002": "su pareja, 700 USD (las 40 camisetas de Gonza)",
        "P-005": "su pareja, 600 USD",
    },
    # Facu, 13/10: para la seña del pedido de China se usaron ~$1.200.000 de ganancias y su
    # pareja puso $420.000 más. La seña (los dos egresos del 22/09) ya está en la caja.
    "aportes": [(date(2026, 9, 22), 420000, "Su pareja, para la seña del pedido de China")],
    # Facu, 13/10: gastos pagados con plata propia (a veces pasa). Entran como aporte y
    # salen como gasto: cuentan en la ganancia y las métricas, la caja no cambia.
    "gastos_con_plata_propia": [
        (date(2026, 10, 13), "Publicidad", 12877, "Publicidad en Marketplace (acumulado)"),
    ],
    # Facu, 13/10: "a día de hoy, con las ventas, en la reserva tenemos 540.000". Aproximado:
    # "no pasa nada si hay diferencia de 20k, el problema serían 100.000".
    "caja_segun_facu": 540000,
    # Facu, 13/10: Cirilo se lleva 9 shorts casual M ($225.000), no 8. En el pedido a China
    # van 2 Liverpool turquesa M y la venta tenía uno solo.
    "items_agregados": [("V-0020", 5010, 1, 25000, "Señada", "Mayorista")],
}

# ---------------------------------------------------------------------------
# El pedido a China actualizado (10-11/10/2026). El proveedor no tenía 8 camisetas de
# P-003 y 20 shorts de P-004; se reemplazaron. Los libros no lo reflejan. Fuentes, en
# C:\Users\facuq\Desktop\RosarioFkits:
#   Pedido_China_SinFaltantes_2026-10-10.xlsx, hoja CAMBIOS   -> las 8 camisetas que salen
#   Para el importador - reemplazos 2026-10-11\1 - Camisetas  -> las 8 que entran (final, con la Lazio)
#   Para el importador - reemplazos 2026-10-11\2 - Shorts     -> los 20 que entran
#   Pedido_Shorts_Proveedor_2026-09-21.xlsx                   -> el pedido real de shorts
# Los 20 shorts que salen no tienen lista con SKU: salen de comparar el pedido real con
# los modelos que el proveedor dijo que no tenía (Argentina negro y blanco, Inter,
# Barcelona 25/26, Chelsea titular, Liverpool sin temporada y Chelsea 25/26 2da): da
# exactamente 20 unidades y 5 con número. verificar_china() lo vuelve a contar contra
# los archivos cada vez que corre.
# SKU nuevos: 4099 en adelante y 5057 en adelante (Facu, 13/10).
# ---------------------------------------------------------------------------
CHINA = {
    "salen": {
        "P-003": [4017, 4018, 4019, 4038, 4049, 4050, 4061],
        "P-004": [5013, 5014, 5015, 5016, 5017, 5018, 5019, 5020, 5021, 5022,
                  5023, 5024, 5027, 5028, 5029, 5030],
    },
    # Los números de los shorts Argentina estaban corridos de modelo en el libro respecto
    # del pedido real del 21/09.
    "estampados": {5017: "N 10", 5018: "N 24", 5034: None, 5037: None, 5038: None},
    # (pedido, sku, modelo, categoria, tipo, talle, estampado, precio, cantidad)
    "nuevos": [
        ("P-003", 4099, "Camiseta Real Madrid 06/07' Suplente (negra)", "Camiseta", "Retro", "L", "BECKHAM 23", 55000, 1),
        ("P-003", 4100, "Camiseta Real Madrid 06/07' Suplente (negra)", "Camiseta", "Retro", "XL", "BECKHAM 23", 55000, 1),
        ("P-003", 4101, "Camiseta Real Madrid 06/07' Suplente (negra)", "Camiseta", "Retro", "L", "RONALDO 9", 55000, 1),
        ("P-003", 4102, "Camiseta Real Madrid 06/07' Suplente (negra)", "Camiseta", "Retro", "XL", "RONALDO 9", 55000, 1),
        ("P-003", 4103, "Camiseta Barcelona 16/17' Titular", "Camiseta", "Retro", "L", "NEYMAR JR 11", 55000, 1),
        ("P-003", 4104, "Camiseta Lazio 00/01' Titular (blanca)", "Camiseta", "Retro", "XL", "VERON 23", 55000, 1),
        ("P-004", 5057, "Short Chelsea 26/27 Suplente (negro y amarillo)", "Short", "Short", "M", None, 45000, 1),
        ("P-004", 5058, "Short Chelsea 26/27 Suplente (negro y amarillo)", "Short", "Short", "L", None, 45000, 4),
        ("P-004", 5059, "Short Chelsea 26/27 Suplente (negro y amarillo)", "Short", "Short", "XL", None, 45000, 4),
        ("P-004", 5060, "Short Argentina Titular (blanco)", "Short", "Short", "L", "N 10", 45000, 2),
        ("P-004", 5061, "Short Argentina Titular (blanco)", "Short", "Short", "XL", "N 24", 45000, 1),
        ("P-004", 5062, "Short Argentina Suplente (azul)", "Short", "Short", "XL", "N 10", 45000, 1),
        ("P-004", 5063, "Short Brasil 26/27 Titular", "Short", "Short", "L", None, 45000, 1),
        ("P-004", 5064, "Short PSG 26/27 Titular", "Short", "Short", "XL", None, 45000, 1),
    ],
    # Reemplazos de un SKU que ya estaba en el pedido: se suma a su cantidad. Las dos
    # Argentina 2006 L con MESSI 19 son "EXTRA" en la hoja CAMBIOS: ya existían (4081, 4079).
    "suman": {("P-003", 4081): 1, ("P-003", 4079): 1, ("P-004", 5033): 1, ("P-004", 5040): 1, ("P-004", 5041): 1,
              ("P-004", 5055): 1, ("P-004", 5056): 1},
    # Facu, 13/10: la Inter amarilla XL que no vino estaba en la reserva de V-0021: se saca.
    "cancelar": [("V-0021", 4038)],
    # Agregadas al pedido por WhatsApp, fuera de los archivos (no son reemplazos). El costo
    # se supone igual al resto del pedido (11 USD por unidad) y se suma a la mercadería.
    "extras": [
        ("P-003", 4105, "Camiseta Real Betis Manga Larga", "Camiseta", "Jugador", "XL", None, 55000, 1),
    ],
}

# Ventas que no están en ningún libro.
VENTAS_NUEVAS = [
    # Facu, 13/10: la Betis manga larga XL ya está pedida para Fabri primo, a $55.000, sin pagar.
    {"fecha": date(2026, 10, 13), "cliente": "Fabri primo", "estado": "Reservada",
     "items": [(4105, 55000)], "notas": "Pedida por WhatsApp; agregada al pedido a China"},
]
CARPETA_PEDIDOS = Path(r"C:\Users\facuq\Desktop\RosarioFkits")

MENSAJE_SIN_MOTIVO = "Migrado: el libro no registraba el motivo"


# ============================================================ utilidades

def sin_tildes(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s or "")) if unicodedata.category(c) != "Mn")


def clave(s):
    """Para comparar nombres: minúsculas, sin tildes, solo letras y números."""
    return re.sub(r"[^a-z0-9]+", " ", sin_tildes(s).lower()).strip()


def fecha(v):
    """El libro mezcla '29/08/2026', números de serie (46300) y vacíos."""
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip()
    for f in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def numero(v):
    """52500, '52500', '$52.500', '27666,67', '27666.67'. None si no se puede."""
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.,-]", "", str(v))
    if not s:
        return None
    sep = max(s.rfind(","), s.rfind("."))
    if sep > -1 and len(s) - sep - 1 in (1, 2):
        s = s[:sep].replace(".", "").replace(",", "") + "." + s[sep + 1:]
    else:
        s = s.replace(".", "").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def col(fila, i):
    return fila[i] if len(fila) > i else ""


def plata(n):
    return "—" if n is None else "$" + f"{n:,.0f}".replace(",", ".")


# ---------------------------------------------------- lógica portada de lib/sheets.ts
# Tiene que dar lo mismo que la web: test_migrar.py lo compara modelo por modelo.

SELECCIONES = [
    "argentina", "brasil", "uruguay", "colombia", "chile", "mexico", "méxico",
    "alemania", "francia", "inglaterra", "españa", "espana", "italia", "portugal",
    "croacia", "canada", "canadá", "gales", "japon", "japón", "holanda",
    "paises bajos", "países bajos", "belgica", "bélgica", "marruecos",
    "estados unidos", "usa", "polonia", "suiza", "dinamarca", "nigeria", "senegal",
]
FIN_DE_EQUIPO = re.compile(
    r"^(\d|\(|home|away|supl|suplente|titular|alternativa|tercera|3ra|entrenamiento|retro|fan|"
    r"player|jugador|edicion|edición|casual|manga|arquero|celeste|negra|negro|blanca|blanco|"
    r"verde|azul|amarilla|roja)", re.I)
NOMBRE_EQUIPO = {
    "boca juniors": "Boca", "river plate": "River", "olympique de marsella": "Marsella",
    "manchester city": "Man City", "manchester united": "Man United", "inter de milan": "Inter",
    "canada": "Canadá", "japon": "Japón", "mexico": "México", "l.a galaxy": "LA Galaxy",
}
NO_SON_EQUIPOS = ["oasis"]


def partir_nombre(nombre):
    m = re.match(r"^(.*?)\s*[-–]\s*Talles?\s+(.+?)\s*$", nombre or "", re.I)
    if m:
        return m.group(1).strip(), m.group(2).strip().upper()
    return (nombre or "").strip(), "ÚNICO"


def clasificar(modelo, categoria_sheet):
    if re.search(r"short", categoria_sheet or "", re.I) or re.match(r"short", modelo, re.I):
        return "Shorts"
    n = sin_tildes(modelo.lower())
    if re.search(r"\bretro\b|aniversario|\b(19|20[01])\d\b", n):
        return "Retro"
    if any(s in n for s in SELECCIONES):
        return "Selecciones"
    return "Clubes"


def equipo_de(modelo):
    palabras = re.sub(r"^(camiseta|short)\s+", "", modelo, flags=re.I).split()
    corte = next((i for i, p in enumerate(palabras) if FIN_DE_EQUIPO.match(p)), None)
    nombre = " ".join(palabras[:corte] if corte is not None else palabras).strip()
    if not nombre:
        return None
    k = sin_tildes(nombre.lower())
    if k in NO_SON_EQUIPOS:
        return None
    return NOMBRE_EQUIPO.get(k, nombre)


# ============================================================ lectura

def leer_libros(desde_json=None):
    if desde_json:
        d = Path(desde_json)
        return (json.loads((d / "viejo.json").read_text(encoding="utf-8")),
                json.loads((d / "nuevo.json").read_text(encoding="utf-8")))
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    # Solo lectura: con este permiso la API no deja escribir aunque el código quisiera.
    cred = service_account.Credentials.from_service_account_file(
        str(CREDENCIALES), scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
    api = build("sheets", "v4", credentials=cred, cache_discovery=False).spreadsheets()
    libros = []
    for lib in (VIEJO, NUEVO):
        hojas = [h["properties"]["title"] for h in api.get(spreadsheetId=lib).execute()["sheets"]]
        libros.append({h: api.values().get(
            spreadsheetId=lib, range=f"'{h}'", valueRenderOption="UNFORMATTED_VALUE",
            dateTimeRenderOption="SERIAL_NUMBER").execute().get("values", []) for h in hojas})
    return libros[0], libros[1]


# ============================================================ armado

class Migracion:
    def __init__(self, viejo, nuevo):
        self.viejo, self.nuevo = viejo, nuevo
        self.obs = defaultdict(list)          # sección del informe -> líneas
        self.skus, self.modelos, self.equipos = {}, {}, set()
        self.proveedores, self.pedidos, self.items = [], {}, []
        self.clientes, self.ventas, self.pagos, self.caja, self.ajustes = {}, {}, [], [], []
        self.viejo_stock = {}                 # sku -> fila del libro viejo

    def nota(self, seccion, texto):
        self.obs[seccion].append(texto)

    # ---------------------------------------------------------------- productos
    def armar_productos(self):
        filas_viejo = self.viejo["STOCK"]
        fila_enc = next(i for i, f in enumerate(filas_viejo) if f and str(f[0]).strip() == "SKU")
        for f in filas_viejo[fila_enc + 1:]:
            if f and str(f[0]).strip().isdigit():
                self.viejo_stock[int(f[0])] = f

        for f in self.nuevo["PRODUCTOS"][1:]:
            if not f or not str(f[0]).strip().isdigit():
                continue
            sku = int(f[0])
            v = self.viejo_stock.get(sku)
            estampado = str(col(f, 3)).strip() or None
            modelo_nuevo = str(col(f, 1)).strip()
            talle_nuevo = str(col(f, 2)).strip().upper()

            if v:
                completo, talle = partir_nombre(str(col(v, 1)))
                modelo = completo
                if estampado and re.search(r"\(\s*" + re.escape(estampado) + r"\s*\)\s*$", completo, re.I):
                    modelo = re.sub(r"\s*\([^)]*\)\s*$", "", completo).strip()
                elif re.search(r"\([^)]*\)\s*$", completo) and not estampado:
                    # "(Guerreros-Azul)" es parte del nombre, no un estampado.
                    pass
                if clave(modelo) != clave(modelo_nuevo):
                    self.nota("nombres", f"SKU {sku}: el viejo dice **{modelo}**, el nuevo "
                                         f"*{modelo_nuevo}*. Se usa el del viejo (es el que ve la web).")
                if talle_nuevo and TALLE_EQUIVALENTE.get(talle_nuevo, talle_nuevo) != talle:
                    self.nota("nombres", f"SKU {sku}: talle {talle} en el viejo y {talle_nuevo} en el nuevo. Se usa {talle}.")
                precio = numero(col(v, 4))
                nuevo_flag = str(col(v, 12)).strip().lower() == "si"
                antes = numero(col(v, 13))
            else:
                modelo, talle = modelo_nuevo, talle_nuevo
                precio = numero(col(f, 7))
                nuevo_flag = str(col(f, 17)).strip().lower() == "si"
                antes = numero(col(f, 18))

            if talle in TALLE_EQUIVALENTE:
                self.nota("automaticas", f"SKU {sku}: talle `{talle}` pasa a `{TALLE_EQUIVALENTE[talle]}`.")
                talle = TALLE_EQUIVALENTE[talle]
            if talle not in TALLES:
                self.nota("bloqueantes", f"SKU {sku}: talle `{talle}` no es válido.")
            if not precio or precio <= 0:
                self.nota("bloqueantes", f"SKU {sku}: sin precio.")
                precio = 1
            if antes is not None and antes <= precio:
                self.nota("automaticas", f"SKU {sku}: Precio Antes {plata(antes)} no es mayor que el precio "
                                         f"{plata(precio)}; se deja sin precio tachado.")
                antes = None

            categoria = str(col(f, 5)).strip() or ("Short" if "short" in modelo.lower() else "Camiseta")
            tipo = str(col(f, 6)).strip() or ("Short" if categoria == "Short" else "Jugador")
            nombre_web = modelo + (f" ({estampado})" if estampado else "")
            linea = clasificar(nombre_web, categoria)
            equipo = equipo_de(nombre_web)
            if equipo:
                self.equipos.add(equipo)

            m = self.modelos.setdefault(modelo, {"categoria": categoria, "tipo": tipo, "linea": linea,
                                                 "equipo": equipo, "nuevo": False, "skus": []})
            if (m["categoria"], m["tipo"], m["linea"]) != (categoria, tipo, linea):
                self.nota("bloqueantes", f"Modelo **{modelo}**: sus talles no coinciden en categoría/tipo/línea "
                                         f"({m['categoria']}/{m['tipo']}/{m['linea']} contra {categoria}/{tipo}/{linea}).")
            m["nuevo"] |= nuevo_flag
            m["skus"].append(sku)
            dorsal_txt = str(col(f, 4)).strip()
            self.skus[sku] = {
                "modelo": modelo, "talle": talle, "estampado": estampado,
                "dorsal": bool(dorsal_txt or estampado), "precio": precio, "precio_antes": antes,
                "estado_pub": str(col(f, 8)).strip() or "Stock",
                "pedido_origen": str(col(f, 9)).strip() or None,
                "nombre_web": f"{nombre_web} - Talle {talle}",
            }

        for sku in self.viejo_stock:
            if sku not in self.skus:
                self.nota("bloqueantes", f"SKU {sku} está en el libro viejo y no en el nuevo.")

        # Dos SKU para la misma unidad (mismo modelo, talle y estampado): la base no lo
        # acepta. Se unen en el más bajo y se pregunta.
        self.fusionados = {}
        grupos = defaultdict(list)
        for sku, s in self.skus.items():
            grupos[(s["modelo"], s["talle"], s["estampado"] or "")].append(sku)
        for (modelo, talle, est), lista in grupos.items():
            if len(lista) > 1:
                queda, *sobran = sorted(lista)
                for s in sobran:
                    self.fusionados[s] = queda
                    del self.skus[s]
                    self.modelos[modelo]["skus"].remove(s)
                    self.viejo_stock.pop(s, None)
                self.nota("automaticas", f"SKU {', '.join(map(str, sorted(lista)))}: {modelo} talle {talle}"
                                         f"{' ' + est if est else ''} estaba cargada {len(lista)} veces; se une en el "
                                         f"{queda} con cantidad {len(lista)} (el pedido a China lleva {len(lista)} iguales).")

    # ---------------------------------------------------------------- pedidos
    def armar_pedidos(self):
        for f in self.nuevo["PEDIDOS"][1:]:
            if not f or not str(f[0]).startswith("P-"):
                continue
            pid = str(f[0]).strip()
            fch = fecha(col(f, 1))
            if isinstance(col(f, 1), (int, float)):
                self.nota("automaticas", f"{pid}: la fecha estaba guardada como número ({col(f, 1)}) = {fch:%d/%m/%Y}.")
            prov = str(col(f, 2)).strip()
            if prov not in self.proveedores:
                self.proveedores.append(prov)
            estado = ESTADOS_PEDIDO.get(clave(col(f, 3)), None)
            if not estado:
                self.nota("bloqueantes", f"{pid}: estado `{col(f, 3)}` desconocido.")
                estado = "Encargado"
            recibido_en = fch if estado == "Recibido" else None
            if pid in CORRECCIONES["pedido_recibido"]:
                recibido_en = CORRECCIONES["pedido_recibido"][pid]
                self.nota("automaticas", f"{pid} figura *{estado}* en el libro nuevo pero ya llegó (Facu, 13/10): "
                                         f"se migra como recibido el {recibido_en:%d/%m/%Y} (fecha aproximada).")
                estado = "Recibido"
            elif estado == "Recibido":
                self.nota("automaticas", f"{pid}: no tiene fecha de llegada; se usa la del pedido ({fch:%d/%m/%Y}).")
            envio = numero(col(f, 6))
            if pid in CORRECCIONES["costo_envio"]:
                self.nota("automaticas", f"{pid}: envío {plata(CORRECCIONES['costo_envio'][pid])} "
                                         f"(el libro decía {plata(envio)} estimado; Facu, 13/10).")
                envio = CORRECCIONES["costo_envio"][pid]
            self.pedidos[pid] = {"fecha": fch, "proveedor": prov, "estado": estado, "recibido_en": recibido_en,
                                 "costo_mercaderia": numero(col(f, 5)), "costo_envio": envio,
                                 "otros": numero(col(f, 7)), "notas": str(col(f, 12)).strip() or None}

        vistos = set()
        for f in self.nuevo["PEDIDO_ITEMS"][1:]:
            if not f or not str(col(f, 0)).startswith("P-"):
                continue
            pid, sku = str(f[0]).strip(), int(f[1])
            if sku in self.fusionados:
                sku = self.fusionados[sku]
                it = next((i for i in self.items if i["pedido_id"] == pid and i["sku"] == sku), None)
                if it:
                    extra = int(numero(col(f, 2)) or 0)
                    it["cantidad"] += extra
                    it["recibida"] += extra if self.pedidos[pid]["estado"] == "Recibido" else 0
                    continue
            if (pid, sku) in vistos:
                self.nota("bloqueantes", f"{pid}: el SKU {sku} aparece dos veces.")
                continue
            vistos.add((pid, sku))
            if sku not in self.skus:
                self.nota("bloqueantes", f"{pid}: el SKU {sku} no está en PRODUCTOS.")
                continue
            cant = int(numero(col(f, 2)) or 0)
            if cant <= 0:
                self.nota("automaticas", f"{pid}, SKU {sku} ({self.skus[sku]['nombre_web']}): figura con cantidad "
                                         f"{cant}; no se carga como comprado (queda el producto, agotado).")
                continue
            ped = self.pedidos[pid]
            recibida = cant if ped["estado"] == "Recibido" else 0
            if (pid, sku) in CORRECCIONES["recibidas"]:
                recibida = CORRECCIONES["recibidas"][(pid, sku)]
                self.nota("automaticas", f"{pid}, SKU {sku}: se pidieron {cant} y llegaron {recibida} (Facu, 13/10).")
            costo = numero(col(f, 3))
            if isinstance(col(f, 3), str) and col(f, 3):
                self.nota("automaticas", f"{pid}, SKU {sku}: el costo estaba como texto (`{col(f, 3)}`).")
            self.items.append({"pedido_id": pid, "sku": sku, "cantidad": cant, "recibida": recibida,
                               "costo_unit": costo, "apellido": str(col(f, 5)).strip() or None,
                               "nota": str(col(f, 6)).strip() or None})

    def aplicar_china(self):
        """El pedido a China actualizado (ver CHINA arriba)."""
        salen = {s for lista in CHINA["salen"].values() for s in lista}
        unidades_salen = defaultdict(int)
        for i in list(self.items):
            if i["sku"] in salen and i["pedido_id"] in CHINA["salen"]:
                unidades_salen[i["pedido_id"]] += i["cantidad"]
                self.items.remove(i)
                self.skus[i["sku"]]["estado_pub"] = "Archivado"
        for sku, est in CHINA["estampados"].items():
            s = self.skus[sku]
            antes = s["estampado"]
            s["estampado"], s["dorsal"] = est, bool(est)
            s["nombre_web"] = s["modelo"] + (f" ({est})" if est else "") + f" - Talle {s['talle']}"
            self.nota("automaticas", f"SKU {sku} {s['modelo']} {s['talle']}: número `{antes or 'sin número'}` → "
                                     f"`{est or 'sin número'}` (estaban corridos respecto del pedido real del 21/09).")
        unidades_entran = defaultdict(int)
        for pid, sku, modelo, cat, tipo, talle, est, precio, cant in CHINA["nuevos"]:
            nombre_web = modelo + (f" ({est})" if est else "")
            linea = clasificar(nombre_web, cat)
            equipo = equipo_de(nombre_web)
            if equipo:
                self.equipos.add(equipo)
            m = self.modelos.setdefault(modelo, {"categoria": cat, "tipo": tipo, "linea": linea,
                                                 "equipo": equipo, "nuevo": False, "skus": []})
            m["skus"].append(sku)
            self.skus[sku] = {"modelo": modelo, "talle": talle, "estampado": est, "dorsal": bool(est),
                              "precio": precio, "precio_antes": None, "estado_pub": "Preventa",
                              "pedido_origen": pid, "nombre_web": f"{nombre_web} - Talle {talle}"}
            self.items.append({"pedido_id": pid, "sku": sku, "cantidad": cant, "recibida": 0, "costo_unit": None,
                               "apellido": est, "nota": "Reemplazo del 10-11/10 (sin stock en el proveedor)"})
            unidades_entran[pid] += cant
        for (pid, sku), cant in CHINA["suman"].items():
            it = next(i for i in self.items if i["pedido_id"] == pid and i["sku"] == sku)
            it["cantidad"] += cant
            unidades_entran[pid] += cant
        for pid, sku, modelo, cat, tipo, talle, est, precio, cant in CHINA["extras"]:
            nombre_web = modelo + (f" ({est})" if est else "")
            equipo = equipo_de(nombre_web)
            if equipo:
                self.equipos.add(equipo)
            m = self.modelos.setdefault(modelo, {"categoria": cat, "tipo": tipo, "linea": clasificar(nombre_web, cat),
                                                 "equipo": equipo, "nuevo": False, "skus": []})
            m["skus"].append(sku)
            self.skus[sku] = {"modelo": modelo, "talle": talle, "estampado": est, "dorsal": bool(est),
                              "precio": precio, "precio_antes": None, "estado_pub": "Preventa",
                              "pedido_origen": pid, "nombre_web": f"{nombre_web} - Talle {talle}"}
            ped = self.pedidos[pid]
            unidades = sum(i["cantidad"] for i in self.items if i["pedido_id"] == pid)
            costo_unit = round((ped["costo_mercaderia"] or 0) / unidades, 2)
            ped["costo_mercaderia"] = round((ped["costo_mercaderia"] or 0) + costo_unit * cant, 2)
            self.items.append({"pedido_id": pid, "sku": sku, "cantidad": cant, "recibida": 0, "costo_unit": None,
                               "apellido": est, "nota": "Agregada por WhatsApp (Facu, 13/10)"})
            self.nota("automaticas", f"{pid}: se agrega {nombre_web} talle {talle} (SKU {sku}), pedida por WhatsApp. "
                                     f"Costo supuesto {plata(costo_unit)} como el resto del pedido; la mercadería pasa a "
                                     f"{plata(ped['costo_mercaderia'])} y {unidades + cant} unidades.")
        for pid in sorted(set(unidades_salen) | set(unidades_entran)):
            self.nota("automaticas", f"{pid}: salen {unidades_salen[pid]} unidades sin stock en el proveedor y entran "
                                     f"{unidades_entran[pid]} de reemplazo (archivos del 10-11/10). Los SKU que salen "
                                     f"quedan archivados.")
            if unidades_salen[pid] != unidades_entran[pid]:
                self.nota("bloqueantes", f"{pid}: salen {unidades_salen[pid]} y entran {unidades_entran[pid]}: no cierra.")
        self.verificar_china()
        if "P-001" in self.pedidos:
            self.nota("automaticas", "P-001: 12 camisetas FAN y 13 shorts; uno de los shorts fue un regalo del "
                                     "proveedor (Facu, 13/10). Los costos de agosto ya eran supuestos: no se cambian.")

    def verificar_china(self):
        """Recuenta contra los archivos que se mandaron al importador."""
        try:
            import openpyxl
        except ImportError:
            self.nota("automaticas", "No se pudo verificar el pedido a China contra los archivos (falta openpyxl).")
            return
        carpeta = CARPETA_PEDIDOS / "Para el importador - reemplazos 2026-10-11"
        for archivo, pid, col_cant in (("1 - Camisetas (8 reemplazos).xlsx", "P-003", 6),
                                       ("2 - Shorts (20 reemplazos).xlsx", "P-004", 5)):
            ruta = carpeta / archivo
            if not ruta.exists():
                self.nota("automaticas", f"No está {ruta.name}: no se verificó {pid}.")
                continue
            ws = openpyxl.load_workbook(ruta, data_only=True).active
            total = sum(int(f[col_cant]) for f in ws.iter_rows(min_row=4, values_only=True)
                        if len(f) > col_cant and isinstance(f[col_cant], (int, float)) and f[col_cant] < 100)
            nuestro = sum(x[8] for x in CHINA["nuevos"] if x[0] == pid) + sum(
                c for (p, _), c in CHINA["suman"].items() if p == pid)
            if total != nuestro:
                self.nota("bloqueantes", f"{pid}: el archivo de reemplazos suma {total} unidades y el script carga {nuestro}.")
        ruta = CARPETA_PEDIDOS / "Pedido_China_SinFaltantes_2026-10-10.xlsx"
        if ruta.exists():
            ws = openpyxl.load_workbook(ruta, data_only=True)["CAMBIOS"]
            skus = {int(f[0]) for f in ws.iter_rows(min_row=3, values_only=True) if isinstance(f[0], (int, float))}
            if skus != set(CHINA["salen"]["P-003"]):
                self.nota("bloqueantes", f"P-003: la hoja CAMBIOS saca {sorted(skus)} y el script {CHINA['salen']['P-003']}.")

    def costo_item(self, pid, sku):
        """Costo real con la importación prorrateada, como v_costo_item."""
        ped = self.pedidos.get(pid)
        it = next((i for i in self.items if i["pedido_id"] == pid and i["sku"] == sku), None)
        if not ped or not it:
            return None
        unidades = sum(i["cantidad"] for i in self.items if i["pedido_id"] == pid)
        base = it["costo_unit"] if it["costo_unit"] is not None else (
            (ped["costo_mercaderia"] or 0) / unidades if unidades else None)
        if base is None:
            return None
        return round(base + ((ped["costo_envio"] or 0) + (ped["otros"] or 0)) / unidades, 2)

    # ---------------------------------------------------------------- clientes y ventas
    def armar_ventas(self):
        for f in self.nuevo["CLIENTES"][1:]:
            if not f or not str(col(f, 0)).startswith("C-"):
                continue
            tel = re.sub(r"\D", "", str(col(f, 2))) or None
            self.clientes[str(f[0]).strip()] = {"nombre": str(col(f, 1)).strip(), "telefono": tel,
                                                "notas": str(col(f, 3)).strip() or None, "tipo": "Minorista"}
        por_nombre = {clave(c["nombre"]): cid for cid, c in self.clientes.items()}

        filas = [f for f in self.nuevo["VENTAS"][1:] if f and str(col(f, 0)).startswith("V-")]
        for f in filas:
            vid = str(f[0]).strip()
            fch = fecha(col(f, 1))
            cli = str(col(f, 2)).strip()
            if cli not in self.clientes:
                self.nota("bloqueantes", f"{vid}: el cliente {cli} no existe.")
                continue
            v = self.ventas.setdefault(vid, {"fecha": fch, "cliente": cli, "canal": None, "cupon": None,
                                             "notas": None, "items": [], "origen": "nuevo"})
            if (v["fecha"], v["cliente"]) != (fch, cli):
                self.nota("bloqueantes", f"{vid}: sus filas tienen distinta fecha o cliente.")
            canal = str(col(f, 11)).strip()
            if canal and clave(canal) in METODOS:
                v["metodo_libro"] = METODOS[clave(canal)]
            elif canal:
                v["canal"] = canal
            v["cupon"] = v["cupon"] or (str(col(f, 12)).strip() or None)
            v["notas"] = v["notas"] or (str(col(f, 14)).strip() or None)
            sku = int(col(f, 3))
            estado = ESTADOS.get(clave(col(f, 10)))
            if not estado:
                self.nota("bloqueantes", f"{vid}: estado `{col(f, 10)}` desconocido.")
                estado = "Entregada"
            motivo = str(col(f, 13)).strip() or None
            if clave(motivo) == "mayorista":
                self.clientes[cli]["tipo"] = "Mayorista"
            costo = numero(col(f, 7))
            if isinstance(col(f, 7), str) and col(f, 7):
                self.nota("automaticas", f"{vid}, SKU {sku}: el costo estaba como texto y rompía la ganancia del libro (#VALUE!).")
            self.agregar_item(vid, sku, int(numero(col(f, 4)) or 1), numero(col(f, 5)), costo, estado, motivo)
        for vid, sku in CHINA["cancelar"]:
            for it in self.ventas.get(vid, {}).get("items", []):
                if it["sku"] == sku and it["estado"] != "Cancelada":
                    it["estado"], it["pedido_id"] = "Cancelada", None
                    self.nota("automaticas", f"{vid}: se saca de la reserva el SKU {sku} "
                                             f"({self.skus[sku]['nombre_web']}): el proveedor no la tenía (Facu, 13/10).")
        for vid, sku, cant, precio, estado, motivo in CORRECCIONES["items_agregados"]:
            self.agregar_item(vid, sku, cant, precio, None, estado, motivo)
            self.nota("automaticas", f"{vid}: se agrega {cant} × {self.skus[sku]['nombre_web']} a {plata(precio)} (Facu, 13/10).")
        n_fechas = sum(1 for f in filas if isinstance(col(f, 1), (int, float)))
        if n_fechas:
            self.nota("automaticas", f"{n_fechas} filas de VENTAS tenían la fecha guardada como número; se convirtieron.")
        if any(v.get("metodo_libro") for v in self.ventas.values()):
            self.nota("automaticas", "La columna `Canal` de VENTAS tiene el medio de pago (Transferencia/Efectivo), "
                                     "no el canal: el canal queda *Sin dato* y el medio sale de PAGOS.")
        for c in self.clientes.values():
            if c["notas"] and "mayorista" in clave(c["notas"]):
                c["tipo"] = "Mayorista"

        # ---- ventas del libro viejo que no están en el nuevo
        enc = next(i for i, f in enumerate(self.viejo["VENTAS"]) if f and str(col(f, 0)).strip() == "ID Venta")
        viejas = []
        for f in self.viejo["VENTAS"][enc + 1:]:
            if not f or not str(col(f, 0)).startswith("V-"):
                continue
            if str(col(f, 12)).strip():
                continue   # copiada desde el panel por el puente (VENTAS!M = "V-0033 #57"): ya está en la base
            viejas.append({"id": str(f[0]).strip(), "fecha": fecha(col(f, 1)),
                           "sku": int(numero(col(f, 2))), "cant": int(numero(col(f, 3)) or 1),
                           "precio": numero(col(f, 5)), "costo": numero(col(f, 4)),
                           "metodo": METODOS.get(clave(col(f, 10))), "cliente": str(col(f, 11)).strip()})
        usados = set()
        pendientes = []
        for vv in viejas:
            pasada = next((vid for vid, v in self.ventas.items()
                           if v["notas"] and f"({vv['id']})" in v["notas"]), None)
            match = None
            for vid, v in self.ventas.items():
                for n, it in enumerate(v["items"]):
                    if (vid, n) in usados or it["sku"] != vv["sku"] or it["precio_unit"] != vv["precio"]:
                        continue
                    if pasada == vid or v["fecha"] == vv["fecha"]:
                        match = (vid, n)
                        break
                if match:
                    break
            if match:
                usados.add(match)
            else:
                pendientes.append(vv)

        siguiente = max(int(v[2:]) for v in self.ventas) + 1
        grupos = defaultdict(list)
        for vv in pendientes:
            grupos[(vv["fecha"], clave(vv["cliente"]))].append(vv)
        for (fch, _), grupo in sorted(grupos.items(), key=lambda g: (g[0][0], g[1][0]["id"])):
            nombre = grupo[0]["cliente"]
            cid = por_nombre.get(clave(nombre)) or por_nombre.get(clave(re.sub(r"[()]", "", nombre)))
            if not cid:
                cid = f"C-{len(self.clientes) + 1:03d}"
                self.clientes[cid] = {"nombre": nombre, "telefono": None, "notas": None, "tipo": "Minorista"}
                por_nombre[clave(nombre)] = cid
                self.nota("clientes", f"{cid} **{nombre}** (nuevo, de la venta {grupo[0]['id']} del libro viejo).")
            vid = f"V-{siguiente:04d}"
            siguiente += 1
            ids = ", ".join(g["id"] for g in grupo)
            self.ventas[vid] = {"fecha": fch, "cliente": cid, "canal": None, "cupon": None,
                                "notas": f"Pasada del libro viejo ({ids})", "items": [], "origen": "viejo"}
            for g in grupo:
                self.agregar_item(vid, g["sku"], g["cant"], g["precio"], g["costo"], "Entregada", None)
            total = sum(g["precio"] * g["cant"] for g in grupo)
            metodo = grupo[0]["metodo"] or "Transferencia"
            self.pagos.append({"id": None, "fecha": fch, "venta_id": vid, "monto": total, "metodo": metodo,
                               "tipo": "Total", "notas": "Cobrada entera al vender (Facu, 13/10: todas por transferencia)"})
            detalle = " + ".join(f"{self.skus[g['sku']]['nombre_web']}" for g in grupo)
            self.nota("automaticas", f"{vid} = {ids} del {fch:%d/%m} (libro viejo): {detalle}, {plata(total)}, "
                                     f"entregada y cobrada por {metodo} (Facu, 13/10).")

    def armar_ventas_nuevas(self):
        por_nombre = {clave(c["nombre"]): cid for cid, c in self.clientes.items()}
        for vn in VENTAS_NUEVAS:
            cid = por_nombre.get(clave(vn["cliente"]))
            if not cid:
                cid = f"C-{len(self.clientes) + 1:03d}"
                self.clientes[cid] = {"nombre": vn["cliente"], "telefono": None, "notas": None, "tipo": "Minorista"}
                por_nombre[clave(vn["cliente"])] = cid
                self.nota("clientes", f"{cid} **{vn['cliente']}** (nuevo, Facu 13/10).")
            vid = f"V-{max(int(v[2:]) for v in self.ventas) + 1:04d}"
            self.ventas[vid] = {"fecha": vn["fecha"], "cliente": cid, "canal": "WhatsApp", "cupon": None,
                                "notas": vn["notas"], "items": [], "origen": "facu"}
            for sku, precio in vn["items"]:
                self.agregar_item(vid, sku, 1, precio, None, vn["estado"], None)
            self.nota("automaticas", f"{vid}: {vn['cliente']}, " + " + ".join(self.skus[s]["nombre_web"] for s, _ in vn["items"])
                      + f", {plata(sum(p for _, p in vn['items']))}, {vn['estado'].lower()} sin pago (Facu, 13/10).")

    def agregar_item(self, vid, sku, cant, precio, costo_libro, estado, motivo):
        sku = self.fusionados.get(sku, sku)
        s = self.skus.get(sku)
        if not s:
            self.nota("bloqueantes", f"{vid}: el SKU {sku} no existe en PRODUCTOS.")
            return
        pid = s["pedido_origen"]
        if pid and not any(i["pedido_id"] == pid and i["sku"] == sku for i in self.items):
            self.nota("automaticas", f"{vid}, SKU {sku}: Pedido_Origen {pid} no tiene ese SKU; se deja sin pedido.")
            pid = None
        costo_real = self.costo_item(pid, sku) if pid else None
        costo = costo_real if (costo_real is not None and pid not in ("P-001", "P-002")) else costo_libro
        if costo is None:
            costo = costo_real
        if costo is None:
            self.nota("bloqueantes", f"{vid}, SKU {sku}: no hay costo.")
            costo = 0
        promo = None
        if clave(motivo) == "mayorista":
            promo, motivo = "Mayorista", None
        if precio is None:
            self.nota("bloqueantes", f"{vid}, SKU {sku}: sin precio.")
            precio = s["precio"]
        if precio < s["precio"] and not promo and not motivo:
            motivo = MENSAJE_SIN_MOTIVO
        self.ventas[vid]["items"].append({
            "sku": sku, "cantidad": cant, "precio_unit": precio, "precio_lista": s["precio"],
            "costo_unit": round(costo, 2),
            # P-001 y P-002 tienen los costos fijos supuestos ($10.000 / $15.000 / $25.500).
            "costo_estimado": pid in ("P-001", "P-002") or pid is None,
            "pedido_id": pid, "estado": estado, "promo": promo, "motivo": motivo})

    # ---------------------------------------------------------------- plata y ajustes
    def armar_plata(self):
        for f in self.nuevo["PAGOS"][1:]:
            if not f or not str(col(f, 0)).startswith("PG-"):
                continue
            vid = str(col(f, 2)).strip()
            if vid not in self.ventas:
                self.nota("bloqueantes", f"{f[0]}: la venta {vid} no existe.")
                continue
            tipo_libro = clave(col(f, 5))
            tipo = TIPOS_PAGO.get(tipo_libro)
            if tipo_libro == "parcial":
                self.nota("automaticas", f"{f[0]}: tipo *Parcial* pasa a *Saldo* (no existe Parcial).")
            cli = str(col(f, 7)).strip()
            if cli and cli != self.ventas[vid]["cliente"]:
                self.nota("bloqueantes", f"{f[0]}: el cliente {cli} no es el de la venta {vid}.")
            self.pagos.insert(len([p for p in self.pagos if p["id"]]), {
                "id": str(f[0]).strip(), "fecha": fecha(col(f, 1)), "venta_id": vid,
                "monto": numero(col(f, 3)), "metodo": METODOS.get(clave(col(f, 4)), "Transferencia"),
                "tipo": tipo or "Total", "notas": str(col(f, 6)).strip() or None})

        for f in self.nuevo["CAJA"][1:]:
            if not f or not str(col(f, 0)).startswith("M-"):
                continue
            ref = str(col(f, 6)).strip()
            cat = next((c for c in CATEGORIAS_CAJA if clave(c) == clave(col(f, 3))), None)
            if not cat:
                self.nota("bloqueantes", f"{f[0]}: categoría `{col(f, 3)}` desconocida.")
                cat = "Otro"
            self.caja.append({"id": str(f[0]).strip(), "fecha": fecha(col(f, 1)),
                              "tipo": "Ingreso" if clave(col(f, 2)) == "ingreso" else "Egreso",
                              "categoria": cat, "monto": numero(col(f, 4)),
                              "metodo": METODOS.get(clave(col(f, 5)), "Transferencia"),
                              "pago_id": ref if ref.startswith("PG-") else None,
                              "pedido_id": ref if ref.startswith("P-") else None,
                              "nota": str(col(f, 7)).strip() or None})
        con_caja = {c["pago_id"] for c in self.caja if c["pago_id"]}
        for p in self.pagos:
            if p["id"] and p["id"] not in con_caja:
                self.nota("bloqueantes", f"{p['id']}: el pago no tiene su movimiento en CAJA.")

        self.caja_extra = []
        for pid, quien in CORRECCIONES["pedidos_pagados_con_aporte"].items():
            ped = self.pedidos.get(pid)
            if not ped:
                continue
            total = ped["costo_mercaderia"] or sum((i["costo_unit"] or 0) * i["cantidad"]
                                                   for i in self.items if i["pedido_id"] == pid)
            total += (ped["costo_envio"] or 0) + (ped["otros"] or 0)
            ya = sum(x["monto"] for x in self.caja if x["pedido_id"] == pid and x["tipo"] == "Egreso")
            if total - ya <= 0:
                continue
            self.caja_extra += [
                {"fecha": ped["fecha"], "tipo": "Ingreso", "categoria": "Aporte", "monto": total - ya,
                 "metodo": "Transferencia", "pago_id": None, "pedido_id": None,
                 "nota": f"Aporte de {quien} para pagar {pid}"},
                {"fecha": ped["fecha"], "tipo": "Egreso", "categoria": "Pago proveedor", "monto": total - ya,
                 "metodo": "Transferencia", "pago_id": None, "pedido_id": pid,
                 "nota": f"{pid} pagado con el aporte de {quien}"}]
            self.nota("automaticas", f"{pid} ({ped['proveedor']}, {plata(total - ya)}): pagado con plata propia "
                                     f"({quien}). Entra como aporte y sale como pago al proveedor el mismo día.")
        for fch, cat, monto, nota in CORRECCIONES["gastos_con_plata_propia"]:
            self.caja_extra += [
                {"fecha": fch, "tipo": "Ingreso", "categoria": "Aporte", "monto": monto, "metodo": "Transferencia",
                 "pago_id": None, "pedido_id": None, "nota": f"Plata propia para: {nota}"},
                {"fecha": fch, "tipo": "Egreso", "categoria": cat, "monto": monto, "metodo": "Transferencia",
                 "pago_id": None, "pedido_id": None, "nota": f"{nota} (pagado con plata propia)"}]
            self.nota("automaticas", f"{cat} {plata(monto)} el {fch:%d/%m}: {nota}, pagado con plata propia (Facu, 13/10).")
        for fch, monto, nota in CORRECCIONES["aportes"]:
            self.caja_extra.append({"fecha": fch, "tipo": "Ingreso", "categoria": "Aporte", "monto": monto,
                                    "metodo": "Transferencia", "pago_id": None, "pedido_id": None, "nota": nota})
            self.nota("automaticas", f"Aporte {plata(monto)} el {fch:%d/%m}: {nota}.")

        for f in self.nuevo["AJUSTES"][1:]:
            if not f or not str(col(f, 1)).strip():
                continue
            sku = int(numero(col(f, 1)))
            self.ajustes.append({"fecha": fecha(col(f, 0)), "sku": self.fusionados.get(sku, sku),
                                 "cantidad": int(numero(col(f, 2))), "motivo": str(col(f, 3)).strip()})

    def armar(self):
        self.armar_productos()
        self.armar_pedidos()
        self.aplicar_china()
        self.armar_ventas()
        self.armar_ventas_nuevas()
        self.armar_plata()
        return self


# ============================================================ carga

TABLAS = ["eventos", "consultas", "stock_diario", "ajustes", "caja_movimientos", "pagos",
          "venta_items", "ventas", "clientes", "pedido_items", "pedidos", "proveedores",
          "skus", "modelos", "equipos"]


def cargar(c, m):
    c.execute("set constraints all deferred")
    c.execute("select set_config('erp.migrando', 'on', true)")
    c.execute("truncate " + ", ".join(TABLAS) + " restart identity")

    eq = {}
    for e in sorted(m.equipos):
        eq[e] = c.execute("insert into equipos (nombre) values (%s) returning id", (e,)).fetchone()[0]
    mod = {}
    for nombre, d in m.modelos.items():
        mod[nombre] = c.execute(
            "insert into modelos (nombre, categoria, tipo, linea, equipo_id, nuevo) values (%s,%s,%s,%s,%s,%s) returning id",
            (nombre, d["categoria"], d["tipo"], d["linea"], eq.get(d["equipo"]), d["nuevo"])).fetchone()[0]
    # executemany en psycopg 3 manda todo en modo pipeline: una ida y vuelta por tabla y no
    # una por fila (contra Oregon, de 5 minutos a segundos).
    cur = c.cursor()
    cur.executemany("insert into skus (sku, modelo_id, talle, estampado, dorsal, precio, precio_antes, estado_pub) "
                    "values (%s,%s,%s,%s,%s,%s,%s,%s)",
                    [(sku, mod[s["modelo"]], s["talle"], s["estampado"], s["dorsal"], s["precio"],
                      s["precio_antes"], s["estado_pub"]) for sku, s in m.skus.items()])
    prov = {p: c.execute("insert into proveedores (nombre) values (%s) returning id", (p,)).fetchone()[0]
            for p in m.proveedores}
    cur.executemany("insert into pedidos (id, fecha, proveedor_id, estado, recibido_en, costo_mercaderia, "
                    "costo_envio, otros, notas) values (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    [(pid, p["fecha"], prov[p["proveedor"]], p["estado"], p["recibido_en"], p["costo_mercaderia"],
                      p["costo_envio"], p["otros"], p["notas"]) for pid, p in m.pedidos.items()])
    cur.executemany("insert into pedido_items (pedido_id, sku, cantidad, cantidad_recibida, costo_unit, apellido, "
                    "nota_estampado) values (%s,%s,%s,%s,%s,%s,%s)",
                    [(i["pedido_id"], i["sku"], i["cantidad"], i["recibida"], i["costo_unit"], i["apellido"],
                      i["nota"]) for i in m.items])
    cur.executemany("insert into clientes (id, nombre, tipo, telefono, notas) values (%s,%s,%s,%s,%s)",
                    [(cid, cl["nombre"], cl["tipo"], cl["telefono"], cl["notas"]) for cid, cl in m.clientes.items()])
    ventas = sorted(m.ventas.items())
    cur.executemany("insert into ventas (id, fecha, cliente_id, canal, cupon, notas) values (%s,%s,%s,%s,%s,%s)",
                    [(vid, v["fecha"], v["cliente"], v["canal"], v["cupon"], v["notas"]) for vid, v in ventas])
    cur.executemany("insert into venta_items (venta_id, sku, cantidad, precio_unit, precio_lista, costo_unit, "
                    "costo_estimado, pedido_id, estado, promo, motivo_descuento, entregado_en) "
                    "values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    [(vid, it["sku"], it["cantidad"], it["precio_unit"], it["precio_lista"], it["costo_unit"],
                      it["costo_estimado"], it["pedido_id"], it["estado"], it["promo"], it["motivo"],
                      datetime.combine(v["fecha"], datetime.min.time()) if it["estado"] == "Entregada" else None)
                     for vid, v in ventas for it in v["items"]])
    # Pagos y caja del libro nuevo, con sus ids, sin que el trigger duplique la caja.
    cur.executemany("insert into pagos (id, fecha, venta_id, monto, metodo, tipo, notas) values (%s,%s,%s,%s,%s,%s,%s)",
                    [(p["id"], p["fecha"], p["venta_id"], p["monto"], p["metodo"], p["tipo"], p["notas"])
                     for p in m.pagos if p["id"]])
    cur.executemany("insert into caja_movimientos (id, fecha, tipo, categoria, monto, metodo, pago_id, pedido_id, nota) "
                    "values (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    [(x["id"], x["fecha"], x["tipo"], x["categoria"], x["monto"], x["metodo"], x["pago_id"],
                      x["pedido_id"], x["nota"]) for x in m.caja])
    cur.executemany("insert into ajustes (fecha, sku, cantidad, motivo) values (%s,%s,%s,%s)",
                    [(a["fecha"], a["sku"], a["cantidad"], a["motivo"]) for a in m.ajustes])

    # Secuencias en el máximo migrado: la próxima venta es la que sigue.
    for seq, tabla, pref in (("venta_seq", "ventas", "V-"), ("pago_seq", "pagos", "PG-"),
                             ("mov_seq", "caja_movimientos", "M-"), ("cliente_seq", "clientes", "C-"),
                             ("pedido_seq", "pedidos", "P-")):
        c.execute(f"select setval('{seq}', greatest(1, coalesce((select max(substr(id, {len(pref) + 1})::int) "
                  f"from {tabla}), 0)), (select count(*) > 0 from {tabla}))")

    # Movimientos agregados por la migración (pedidos pagados antes de la caja), con el
    # número que sigue.
    for x in m.caja_extra:
        c.execute("insert into caja_movimientos (fecha, tipo, categoria, monto, metodo, pedido_id, nota) "
                  "values (%s,%s,%s,%s,%s,%s,%s)",
                  (x["fecha"], x["tipo"], x["categoria"], x["monto"], x["metodo"], x["pedido_id"], x["nota"]))
    # Los pagos supuestos de las ventas del libro viejo: ahora sí con el trigger, que
    # genera su caja con el número que sigue.
    c.execute("select set_config('erp.migrando', 'off', true)")
    for p in m.pagos:
        if not p["id"]:
            c.execute("insert into pagos (fecha, venta_id, monto, metodo, tipo, notas) values (%s,%s,%s,%s,%s,%s)",
                      (p["fecha"], p["venta_id"], p["monto"], p["metodo"], p["tipo"], p["notas"]))
    # La carga no es un "evento": la auditoría arranca después de la migración.
    c.execute("delete from eventos")


# ============================================================ controles e informe

def controlar(c, m):
    filas = c.execute("select sku, esperado, recibido, entregado, comprometido, ajustes, fisico, disponible "
                      "from v_stock order by sku").fetchall()
    stock = {f[0]: f for f in filas}
    difs = []
    for sku, v in sorted(m.viejo_stock.items()):
        if sku not in stock:
            continue
        disp_viejo = int(numero(col(v, 8)) or 0)
        _, esp, rec, ent, comp, aj, fis, disp = stock[sku]
        if disp != disp_viejo:
            difs.append(f"SKU {sku} {m.skus[sku]['nombre_web']}: la web muestra **{disp_viejo}**, el ERP calcula "
                        f"**{disp}** (comprado {esp}, entregado {ent}, comprometido {comp}, ajustes {aj}).")
    negativos = [f"SKU {f[0]}: disponible {f[7]}, en casa {f[6]}" for f in filas if f[7] < 0 or f[6] < 0]

    tab = dict(zip([d.name for d in c.execute("select * from v_tablero").description],
                   c.execute("select * from v_tablero").fetchone()))
    libro = {}
    for f in m.nuevo.get("TABLERO", []):
        if len(f) > 1 and isinstance(f[1], (int, float)):
            libro[clave(f[0])] = f[1]
    comparacion = [
        ("Saldo de caja", libro.get("saldo de caja"), tab["saldo_caja"]),
        ("Facturado", libro.get("facturado"), tab["facturado"]),
        ("Cobrado", libro.get("cobrado"), tab["cobrado"]),
        ("A cobrar", libro.get("a cobrar"), tab["a_cobrar"]),
        ("Deuda con proveedores", libro.get("deuda con el importador"), tab["deuda_proveedores"]),
        ("Ganancia", libro.get("ganancia"), tab["ganancia"]),
        ("Caja según Facu (13/10, aproximada)", CORRECCIONES["caja_segun_facu"], tab["saldo_caja"]),
    ]
    conteos = {t: c.execute(f"select count(*) from {t}").fetchone()[0]
               for t in ("modelos", "skus", "pedidos", "pedido_items", "clientes", "ventas", "venta_items",
                         "pagos", "caja_movimientos", "ajustes")}
    return difs, negativos, comparacion, conteos


def informe(m, difs, negativos, comparacion, conteos, error_final, modo, destino):
    hoy = datetime.now()
    L = [f"# Informe de migración — {hoy:%d/%m/%Y %H:%M}", "",
         f"Modo `{modo}` contra la base **{destino}**. Los libros se leyeron con permiso de solo lectura.", ""]
    if error_final:
        L += ["## ⛔ La base rechazó la carga", "", f"`{error_final}`", ""]
    L += ["## Resumen", "", "| | |", "|---|---|"] + [f"| {k} | {v} |" for k, v in conteos.items()] + [""]

    def seccion(titulo, lineas, explicacion=None):
        nonlocal L
        if not lineas:
            return
        L += [f"## {titulo}", ""]
        if explicacion:
            L += [explicacion, ""]
        L += [f"- {x}" for x in lineas] + [""]

    seccion("⛔ Bloqueantes", m.obs["bloqueantes"] + negativos,
            "Sin resolver esto no se puede hacer el corte.")
    seccion("❓ Para que Facu confirme", m.obs["confirmar"],
            "Supuestos que hizo el script. Si alguno está mal, se corrige acá y se vuelve a correr.")
    seccion("📦 Stock: la web (libro viejo) contra el ERP", difs,
            "Si el ERP da distinto que la web, uno de los dos está mal. El ERP lo calcula de pedidos, ventas y ajustes.")
    L += ["## 💰 Tablero: libro nuevo contra el ERP", "",
          "El libro nuevo dejó de cargarse en la unificación: las diferencias tienen que salir de las "
          "ventas del libro viejo que se suman (y sus pagos supuestos).", "",
          "| | Libro nuevo | ERP | Diferencia |", "|---|---|---|---|"]
    for nombre, lib, erp in comparacion:
        erp = float(erp) if erp is not None else None
        if lib is None:
            L.append(f"| {nombre} | (el libro da error) | {plata(erp)} | — |")
            continue
        dif = erp - lib
        L.append(f"| {nombre} | {plata(lib)} | {plata(erp)} | {plata(dif) if round(dif) else 'igual'} |")
    L.append("")
    seccion("✏️ Nombres distintos entre los libros", m.obs["nombres"])
    seccion("👤 Clientes que se crean", m.obs["clientes"])
    seccion("🔧 Corregido automáticamente", sorted(set(m.obs["automaticas"]), key=m.obs["automaticas"].index))
    return "\n".join(L) + "\n"


# ============================================================ main

def url_real():
    for linea in (RAIZ / ".env.local").read_text(encoding="utf-8").splitlines():
        if linea.startswith("SUPABASE_DB_URL="):
            return linea.split("=", 1)[1].strip().strip('"').strip("'")
    sys.exit("Falta SUPABASE_DB_URL en .env.local")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    modo = ap.add_mutually_exclusive_group(required=True)
    modo.add_argument("--revisar", action="store_true", help="carga, controla y deshace: solo el informe")
    modo.add_argument("--ensayo", action="store_true", help="carga y confirma")
    modo.add_argument("--final", action="store_true", help="el corte (pide confirmación)")
    ap.add_argument("--destino", choices=["local", "real"], default="local")
    ap.add_argument("--desde-json", help="carpeta con viejo.json y nuevo.json (para probar sin Google)")
    a = ap.parse_args()
    nombre_modo = "revisar" if a.revisar else "ensayo" if a.ensayo else "final"

    import psycopg
    if a.destino == "real" and not a.revisar:
        if input(f"Esto BORRA y vuelve a cargar la base REAL ({nombre_modo}). Escribí 'si' para seguir: ").strip() != "si":
            sys.exit("Cancelado.")
    if a.final and input("Es el CORTE. ¿Los libros ya están congelados? Escribí 'corte': ").strip() != "corte":
        sys.exit("Cancelado.")

    viejo, nuevo = leer_libros(a.desde_json)
    m = Migracion(viejo, nuevo).armar()

    error_final = None
    with psycopg.connect(url_real() if a.destino == "real" else LOCAL) as c:
        # Desde el 17/10 el panel carga ventas reales en la base: truncar y recargar desde
        # los libros las borraría (con sus cobros y gastos). El puente las copia al libro
        # viejo solo para la web; el corte ya no pasa por acá.
        if a.destino == "real" and not a.revisar:
            n = c.execute("select count(*) from ventas where creado_por is not null").fetchone()[0]
            if n:
                sys.exit(f"La base real tiene {n} ventas cargadas desde el panel: recargarla las borra. "
                         "No se migra más sobre la real (ver herramientas/backups/LEEME.md, Puente).")
        cargar(c, m)
        try:
            c.execute("set constraints all immediate")   # que salten ya las reglas de stock
        except psycopg.Error as e:
            error_final = e.diag.message_primary
            c.rollback()
            # Para el informe igual: cargar sin las reglas al final.
            cargar(c, m)
        difs, negativos, comparacion, conteos = controlar(c, m)
        if a.revisar or error_final:
            c.rollback()
        else:
            c.commit()

    dest = AQUI / "informes"
    dest.mkdir(exist_ok=True)
    archivo = dest / f"informe-{datetime.now():%Y-%m-%d-%H%M}-{nombre_modo}-{a.destino}.md"
    archivo.write_text(informe(m, difs, negativos, comparacion, conteos, error_final, nombre_modo, a.destino),
                       encoding="utf-8")
    print(f"Informe: {archivo}")
    print("Resultado:", "RECHAZADA (" + error_final + ")" if error_final else
          ("deshecha (revisar)" if a.revisar else "CONFIRMADA"))
    print(f"Bloqueantes: {len(m.obs['bloqueantes']) + len(negativos)}  Para confirmar: {len(m.obs['confirmar'])}  "
          f"Diferencias de stock con la web: {len(difs)}")


if __name__ == "__main__":
    main()
